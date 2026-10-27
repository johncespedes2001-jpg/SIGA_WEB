from __future__ import annotations

from pathlib import Path
from textwrap import wrap
from zipfile import ZIP_DEFLATED, ZipFile
import shutil
import tempfile
import xml.etree.ElementTree as ET

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\SIGA_WEB\SIGA-Web")
WORK = ROOT / "informe_trabajo"
ASSETS = WORK / "assets"
SOURCE = Path(r"C:\Users\JAIRO\Downloads\FORMATO DE INFORME.docx")
OUTPUT = ROOT / "Informe_SIGA_hasta_punto_8_evidencias_actualizadas.docx"

NAVY = "#17365D"
BLUE = "#285EA8"
LIGHT_BLUE = "#EAF2FB"
PALE_BLUE = "#F4F8FC"
GRAY = "#D9E2F3"
DARK = "#16243D"
TEAL = "#178A8A"
GREEN = "#2E7D4F"
WHITE = "#FFFFFF"
BLACK = "#000000"


def font(name: str, size: int, bold: bool = False):
    candidates = [
        Path(rf"C:\Windows\Fonts\{name}.ttf"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


FONT_TITLE = font("arialbd", 38, True)
FONT_SUBTITLE = font("arialbd", 24, True)
FONT_BODY = font("arial", 19)
FONT_SMALL = font("arial", 16)


def centered_text(draw: ImageDraw.ImageDraw, box, text, text_font, fill=DARK):
    x1, y1, x2, y2 = box
    lines = []
    for paragraph in text.split("\n"):
        lines.extend(wrap(paragraph, width=max(18, int((x2 - x1) / (text_font.size * 0.62)))))
    line_height = text_font.size + 7
    y = y1 + ((y2 - y1) - line_height * len(lines)) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=text_font)
        x = x1 + ((x2 - x1) - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=text_font, fill=fill)
        y += line_height


def arrow(draw, start, end, color=BLUE, width=5):
    draw.line([start, end], fill=color, width=width)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        points = [(ex, ey), (ex - 16 * direction, ey - 9), (ex - 16 * direction, ey + 9)]
    else:
        direction = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 9, ey - 16 * direction), (ex + 9, ey - 16 * direction)]
    draw.polygon(points, fill=color)


def make_architecture(path: Path):
    image = Image.new("RGB", (1700, 900), WHITE)
    draw = ImageDraw.Draw(image)
    draw.text((70, 45), "Arquitectura final de SIGA Web", font=FONT_TITLE, fill=DARK)
    draw.text((72, 100), "Flujo de la interfaz, la API y la persistencia institucional", font=FONT_BODY, fill="#4B627F")

    boxes = [
        ((70, 210, 380, 640), LIGHT_BLUE, "Usuarios y navegador", "Administrador\nCoordinador\nDocente y Tutor\nEstudiante\nApoderado\nVigilante"),
        ((500, 210, 800, 640), "#E9F7F5", "Frontend", "HTML5\nCSS3 responsive\nJavaScript\n64 vistas por rol"),
        ((920, 150, 1250, 700), "#EEF0FF", "Backend y API", "Node.js 18+\nExpress 5.3\n149 rutas API\nAutenticación\nCSRF y roles\nReglas de negocio"),
        ((1370, 210, 1630, 640), "#F2F8ED", "Datos", "SQL Server 2025\nBase SIGA\n28 tablas\nProcedimiento de\nsincronización"),
    ]
    for rect, fill, title, body in boxes:
        draw.rounded_rectangle(rect, radius=28, fill=fill, outline="#B5C8E5", width=4)
        x1, y1, x2, y2 = rect
        draw.rounded_rectangle((x1 + 24, y1 + 24, x2 - 24, y1 + 94), radius=18, fill=NAVY)
        centered_text(draw, (x1 + 35, y1 + 30, x2 - 35, y1 + 88), title, FONT_SUBTITLE, WHITE)
        centered_text(draw, (x1 + 35, y1 + 120, x2 - 35, y2 - 35), body, FONT_BODY)

    arrow(draw, (380, 425), (500, 425))
    arrow(draw, (800, 425), (920, 425))
    arrow(draw, (1250, 425), (1370, 425))
    draw.text((425, 378), "HTTPS / JSON", font=FONT_SMALL, fill=BLUE)
    draw.text((825, 378), "API REST", font=FONT_SMALL, fill=BLUE)
    draw.text((1267, 378), "mssql", font=FONT_SMALL, fill=BLUE)

    draw.rounded_rectangle((500, 745, 1250, 835), radius=22, fill="#F7F9FC", outline="#C6D3E4", width=3)
    centered_text(draw, (520, 755, 1230, 825), "ExcelJS · Multer · PDFKit · QRCode · compresión · ETag · pruebas con node:test", FONT_BODY)
    image.save(path, quality=95)


def make_erd(path: Path):
    image = Image.new("RGB", (1800, 1180), WHITE)
    draw = ImageDraw.Draw(image)
    draw.text((65, 40), "Modelo entidad relación principal de SIGA", font=FONT_TITLE, fill=DARK)
    draw.text((67, 94), "Vista simplificada de las entidades centrales y sus relaciones", font=FONT_BODY, fill="#4B627F")

    nodes = {
        "Programas": (70, 180, 390, 330, ["ProgramaId PK", "CodigoPublico", "Nombre", "Estado"]),
        "Usuarios": (520, 160, 880, 340, ["UsuarioId PK", "ProgramaId FK", "NombreUsuario", "ContrasenaHash", "Estado"]),
        "Roles": (1010, 180, 1290, 330, ["RolId PK", "Clave", "Nombre"]),
        "UsuariosRoles": (1430, 180, 1730, 340, ["UsuarioRolId PK", "UsuarioId FK", "RolId FK", "EsPrincipal"]),
        "Estudiantes": (70, 470, 430, 670, ["EstudianteId PK", "UsuarioId FK", "ProgramaId FK", "TutorUsuarioId FK", "Semestre", "Grupo"]),
        "Personal": (520, 480, 860, 650, ["PersonalId PK", "UsuarioId FK", "ProgramaId FK", "CodigoEmpleado", "Cargo"]),
        "Clases": (1010, 450, 1370, 680, ["ClaseId PK", "ProgramaId FK", "DocenteUsuarioId FK", "Curso", "Semestre y grupo", "Día y horario", "Aula y modalidad"]),
        "MatriculasClase": (1450, 480, 1740, 660, ["MatriculaClaseId PK", "ClaseId FK", "EstudianteId FK", "Estado"]),
        "RegistrosAsistencia": (70, 820, 430, 1090, ["AsistenciaId PK", "ClaseId FK", "EstudianteId FK", "Fecha", "Estado", "RegistradoPor FK"]),
        "SolicitudesSalida": (520, 820, 900, 1090, ["SolicitudSalidaId PK", "EstudianteId FK", "ClaseId FK", "Estado", "CodigoQr", "AutorizadoPor FK"]),
        "Justificaciones": (1010, 830, 1370, 1080, ["JustificacionId PK", "EstudianteId FK", "Estado", "Evidencia", "RevisadoPor FK"]),
        "ValidacionesSalida": (1450, 830, 1740, 1080, ["ValidacionSalidaId PK", "SolicitudSalidaId FK", "VigilanteUsuarioId FK", "FechaValidacion"]),
    }

    centers = {name: ((values[0] + values[2]) // 2, (values[1] + values[3]) // 2) for name, values in nodes.items()}
    relations = [
        ("Programas", "Usuarios"), ("Usuarios", "UsuariosRoles"), ("Roles", "UsuariosRoles"),
        ("Programas", "Estudiantes"), ("Usuarios", "Estudiantes"), ("Usuarios", "Personal"),
        ("Programas", "Clases"), ("Usuarios", "Clases"), ("Clases", "MatriculasClase"),
        ("Estudiantes", "MatriculasClase"), ("Clases", "RegistrosAsistencia"),
        ("Estudiantes", "RegistrosAsistencia"), ("Estudiantes", "SolicitudesSalida"),
        ("SolicitudesSalida", "ValidacionesSalida"), ("Estudiantes", "Justificaciones"),
    ]
    for source, target in relations:
        sx, sy = centers[source]
        tx, ty = centers[target]
        draw.line((sx, sy, tx, ty), fill="#9AB0CF", width=3)

    # Las entidades se dibujan después de las relaciones para que ninguna línea
    # atraviese los títulos o los campos de las tablas.
    for name, (x1, y1, x2, y2, fields) in nodes.items():
        draw.rounded_rectangle((x1, y1, x2, y2), radius=20, fill=PALE_BLUE, outline="#9EB7D8", width=3)
        draw.rounded_rectangle((x1, y1, x2, y1 + 54), radius=20, fill=NAVY)
        draw.rectangle((x1, y1 + 35, x2, y1 + 54), fill=NAVY)
        centered_text(draw, (x1 + 12, y1 + 5, x2 - 12, y1 + 50), name, FONT_SUBTITLE, WHITE)
        y = y1 + 66
        for field in fields:
            draw.text((x1 + 20, y), field, font=FONT_SMALL, fill=DARK)
            y += 27
    image.save(path, quality=95)


def set_run_font(run, name="Arial", size=11, bold=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))


def replace_paragraph(paragraph, text, style=None, size=11, bold=False, align=None):
    paragraph.clear()
    if style:
        try:
            paragraph.style = style
        except KeyError:
            pass
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.08
    return paragraph


def clear_paragraph(paragraph):
    paragraph.clear()
    paragraph.paragraph_format.space_after = Pt(0)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.lstrip("#"))


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def style_table(table, font_size=9):
    table.autofit = True
    set_table_borders(table)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, NAVY if row_index == 0 else (PALE_BLUE if row_index % 2 == 0 else WHITE))
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    set_run_font(run, size=font_size, bold=(row_index == 0), color=(WHITE if row_index == 0 else BLACK))


def mark_table_header(table):
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "true")


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)


def add_alt_text(shape, title, description):
    doc_pr = shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def insert_paragraph_after(anchor, text="", style=None):
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    paragraph = Paragraph(new_p, anchor._parent)
    if style:
        try:
            paragraph.style = style
        except KeyError:
            pass
    if text:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def insert_paragraph_after_table(table, text="", style=None):
    new_p = OxmlElement("w:p")
    table._tbl.addnext(new_p)
    from docx.text.paragraph import Paragraph
    paragraph = Paragraph(new_p, table._parent)
    if style:
        try:
            paragraph.style = style
        except KeyError:
            pass
    if text:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def insert_picture_after(anchor, image_path, width_inches, title, description):
    paragraph = insert_paragraph_after(anchor)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(4)
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(width_inches))
    add_alt_text(shape, title, description)
    return paragraph


def insert_caption_after(anchor, text):
    paragraph = insert_paragraph_after(anchor)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_run_font(run, size=9, bold=False, color="4A5B73")
    run.italic = True
    return paragraph


def insert_requirement_evidence(anchor, code, title, image_path, caption, page_break=False):
    heading = insert_paragraph_after(anchor)
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(5)
    heading.paragraph_format.keep_with_next = True
    if page_break:
        heading.paragraph_format.page_break_before = True
    code_run = heading.add_run(f"{code}. ")
    set_run_font(code_run, size=11, bold=True, color=NAVY)
    title_run = heading.add_run(title)
    set_run_font(title_run, size=11, bold=True, color=DARK)

    picture = insert_picture_after(
        heading,
        image_path,
        7.05,
        f"Evidencia visual {code}",
        caption,
    )
    picture.paragraph_format.keep_with_next = True
    return insert_caption_after(picture, caption)


def paragraph_by_text(document, text):
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise KeyError(text)


def paragraph_after_heading(document, heading, offset=1):
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() == heading:
            return paragraphs[index + offset]
    raise KeyError(heading)


def fill_table(table, rows):
    for row_index, values in enumerate(rows, start=1):
        for col_index, value in enumerate(values):
            set_cell_text(table.rows[row_index].cells[col_index], value, bold=(col_index == 0 and len(values) == 2))
    style_table(table)


def patch_missing_alt_text(docx_path: Path):
    namespaces = {
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    }
    with tempfile.TemporaryDirectory(dir=str(WORK)) as temp_dir:
        patched = Path(temp_dir) / docx_path.name
        with ZipFile(docx_path, "r") as source, ZipFile(patched, "w", ZIP_DEFLATED) as target:
            for info in source.infolist():
                payload = source.read(info.filename)
                if info.filename.startswith("word/") and info.filename.endswith(".xml"):
                    try:
                        root = ET.fromstring(payload)
                        changed = False
                        for xpath in (".//wp:docPr", ".//pic:cNvPr"):
                            for element in root.findall(xpath, namespaces):
                                if not element.get("descr") and not element.get("title"):
                                    element.set("title", "Elemento gráfico de la plantilla")
                                    element.set("descr", "Elemento gráfico decorativo que organiza visualmente el informe.")
                                    changed = True
                        if changed:
                            payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                    except ET.ParseError:
                        pass
                target.writestr(info, payload)
        shutil.copyfile(patched, docx_path)


def build():
    WORK.mkdir(parents=True, exist_ok=True)
    architecture = ASSETS / "arquitectura-siga-redisenada.png"
    erd = ASSETS / "modelo-er-siga.png"
    make_architecture(architecture)
    make_erd(erd)

    document = Document(str(SOURCE))
    document.core_properties.title = "Informe final del proyecto SIGA Web"
    document.core_properties.subject = "Secciones 1 a 8 del informe final de Taller de Programación Web"
    document.core_properties.author = "Jairo Céspedes"
    document.core_properties.keywords = "SIGA, Node.js, Express, SQL Server, gestión de asistencia"

    # 1. Datos generales
    fill_table(document.tables[3], [
        ["Nombre del proyecto", "SIGA Web - Sistema Integrado de Gestión de Asistencia"],
        ["Integrantes", "Jairo Céspedes"],
        ["Coordinador/a", "Jairo Céspedes"],
        ["Usuarios objetivo", "Administrador, Coordinador, Docente, Tutor, Estudiante, Padre o apoderado y Vigilante"],
        ["Docente", "Ing. Brigitte Susan Flores Zapata"],
        ["Periodo", "2026-II"],
        ["Fecha de presentación", "22/10/2026"],
    ])

    # 2. Resumen ejecutivo
    replace_paragraph(paragraph_after_heading(document, "RESUMEN EJECUTIVO", 1),
        "SIGA Web atiende la necesidad de registrar y consultar asistencia, horarios, permisos de salida, justificaciones y casos de permanencia sin depender de formatos dispersos. La solución organiza la información académica del IESTP Contralmirante Manuel Villar Olivera y asigna a cada perfil únicamente las funciones y los datos que necesita.")
    replace_paragraph(paragraph_after_heading(document, "RESUMEN EJECUTIVO", 2),
        "La aplicación se desarrolló con HTML5, CSS3 y JavaScript en el frontend; Node.js y Express en el backend; y Microsoft SQL Server para la persistencia institucional. Incluye siete perfiles, 64 vistas HTML y 149 rutas de API. Al cierre del proyecto se encuentran operativos el acceso por roles, la gestión de estudiantes y personal, las clases y horarios, la asistencia, los permisos con código QR, las justificaciones, los reportes y la carga masiva mediante Excel.")
    for offset in range(3, 9):
        clear_paragraph(paragraph_after_heading(document, "RESUMEN EJECUTIVO", offset))

    # 3. Problema, necesidad y justificación
    replace_paragraph(paragraph_after_heading(document, "Problema identificado", 1),
        "La institución necesitaba centralizar procesos que involucraban a diferentes responsables: matrícula y asignación académica, control diario de asistencia, permisos de salida, seguimiento tutorial y comunicación con las familias. Cuando estos registros se administran por separado, se dificulta conocer la situación actual de cada estudiante, se duplican datos y aumenta el riesgo de que una persona consulte o modifique información fuera de su responsabilidad.")
    replace_paragraph(paragraph_after_heading(document, "Problema identificado", 2),
        "También era necesario que los docentes visualizaran solo sus clases, los estudiantes solo su horario y sus registros, los coordinadores únicamente su programa de estudios y vigilancia la información mínima para validar una salida.")
    clear_paragraph(paragraph_after_heading(document, "Problema identificado", 3))

    replace_paragraph(paragraph_after_heading(document, "Justificación", 1),
        "La solución es pertinente porque reúne los procesos principales en una sola plataforma, conserva trazabilidad y reduce consultas manuales entre áreas. El control por rol protege el alcance de la información, mientras que la base de datos permite mantener relaciones consistentes entre programas, usuarios, clases, matrículas, asistencia y solicitudes.")
    replace_paragraph(paragraph_after_heading(document, "Justificación", 2),
        "SIGA también facilita la toma de decisiones mediante indicadores, filtros y exportaciones. El uso de una interfaz responsive permite operar desde computadoras y dispositivos móviles en coordinación académica, aula, tutoría y puerta de ingreso.")
    clear_paragraph(paragraph_after_heading(document, "Justificación", 3))

    replace_paragraph(paragraph_after_heading(document, "Usuarios del sistema", 1),
        "La plataforma separa responsabilidades para evitar accesos innecesarios y mantener un flujo institucional verificable.")
    fill_table(document.tables[4], [
        ["Administrador", "Gestiona programas, personal, vigilantes, vínculos familiares, reportes y configuración general."],
        ["Coordinador, Docente y Tutor", "Organizan estudiantes, clases y horarios; registran asistencia; revisan permisos, justificaciones y permanencia según su asignación."],
        ["Estudiante", "Consulta su horario y asistencia, solicita salidas, presenta justificaciones y revisa avisos y seguimiento autorizado."],
        ["Padre o apoderado y Vigilante", "El apoderado consulta a sus estudiantes vinculados; vigilancia valida ingresos y salidas mediante códigos autorizados."],
    ])
    for offset in range(2, 7):
        try:
            clear_paragraph(paragraph_after_heading(document, "Usuarios del sistema", offset))
        except (KeyError, IndexError):
            break

    # 4. Objetivos
    replace_paragraph(paragraph_after_heading(document, "Objetivo general", 1),
        "Desarrollar e integrar una plataforma web multiperfil para gestionar asistencia, horarios, permisos, permanencia y comunicación académica del IESTP Contralmirante Manuel Villar Olivera, con una interfaz responsive, control de acceso por roles y persistencia en Microsoft SQL Server.")
    clear_paragraph(paragraph_after_heading(document, "Objetivo general", 2))
    clear_paragraph(paragraph_after_heading(document, "Objetivo general", 3))
    clear_paragraph(paragraph_after_heading(document, "Objetivo general", 4))
    objectives = [
        "Implementar autenticación segura, sesiones protegidas y autorización por rol y ámbito académico.",
        "Digitalizar el registro de estudiantes, personal, clases, horarios, matrículas y asistencia.",
        "Integrar solicitudes de salida, códigos QR, justificaciones, tutoría y permanencia con trazabilidad.",
        "Facilitar la consulta y el análisis mediante paneles, filtros, reportes y exportaciones a Excel y PDF.",
    ]
    for offset, text in enumerate(objectives, 1):
        replace_paragraph(paragraph_after_heading(document, "Objetivos específicos", offset), text, style="List Paragraph")

    # 5. Requerimientos
    replace_paragraph(paragraph_after_heading(document, "REQUERIMIENTOS FUNCIONALES IMPLEMENTADOS", 1),
        "Los requerimientos se verificaron contra las rutas del servidor, las vistas por rol y las pruebas automatizadas disponibles al cierre de la Semana 16.")
    req_rows = [
        ["RF01", "Autenticar usuarios y limitar cada sesión por rol, estado y ámbito académico.", "Alta", "Implementado", "Evidencias RF01-A y RF01-B: cambio Coordinador/Docente y menús restringidos."],
        ["RF02", "Gestionar programas, personal, vigilantes y estudiantes, incluida carga masiva Excel.", "Alta", "Implementado", "Evidencia RF02: directorio, filtros, plantilla, carga masiva y registro."],
        ["RF03", "Administrar clases, matrículas y horarios visibles solo para docentes y estudiantes asignados.", "Alta", "Implementado", "Evidencias RF03-A y RF03-B: clases, matrículas y nuevo módulo Asignar horarios."],
        ["RF04", "Registrar, corregir y consultar asistencia con reportes exportables.", "Alta", "Implementado", "Evidencias RF04-A y RF04-B: registro de asistencia y reportes Excel/PDF."],
        ["RF05", "Gestionar solicitudes de salida y validar un código QR de uso único.", "Media", "Implementado", "Evidencia RF05: solicitud, autorización de Coordinación y validación por Vigilancia."],
        ["RF06", "Gestionar justificaciones, tutoría, permanencia, avisos y vínculos familiares.", "Media", "Implementado", "Evidencias RF06-A y RF06-B: justificaciones e incidencias con seguimiento."],
    ]
    for i, row in enumerate(req_rows, 1):
        for j, value in enumerate(row):
            set_cell_text(document.tables[5].rows[i].cells[j], value, size=8.5)
    style_table(document.tables[5], font_size=8.5)

    evidence_anchor = insert_paragraph_after_table(document.tables[5])
    replace_paragraph(
        evidence_anchor,
        "Evidencias visuales de los requerimientos",
        style="Heading 2",
        size=14,
        bold=True,
    )
    evidence_anchor.paragraph_format.page_break_before = True
    evidence_anchor.paragraph_format.keep_with_next = True
    evidence_anchor = insert_paragraph_after(
        evidence_anchor,
        "Las siguientes capturas relacionan cada requerimiento con una pantalla funcional del sistema. Los nombres de usuario y los datos visibles corresponden al entorno local de demostración.",
    )
    evidence_anchor.paragraph_format.space_after = Pt(8)

    evidence_anchor = insert_requirement_evidence(
        evidence_anchor,
        "RF01-A",
        "Sesión autenticada en modo Coordinador",
        ASSETS / "actual-rf01-coordinador.jpg",
        "Evidencia RF01-A. Captura actual del 22/10/2026: sesión de Ana Rodríguez Silva en modo Coordinador, con cierre de sesión, selector de rol y acceso limitado al programa Mecatrónica Automotriz.",
    )
    evidence_anchor = insert_requirement_evidence(
        evidence_anchor,
        "RF01-B",
        "Cambio autorizado al modo Docente",
        ASSETS / "actual-rf01-docente.jpg",
        "Evidencia RF01-B. La misma cuenta cambia a modo Docente sin reiniciar sesión y recibe un menú reducido a horario, asistencia, estudiantes, boletas, reportes y perfil.",
        page_break=True,
    )
    evidence_anchor = insert_requirement_evidence(
        evidence_anchor,
        "RF02",
        "Gestión de estudiantes y carga masiva",
        ASSETS / "actual-rf02-estudiantes.jpg",
        "Evidencia RF02. Captura actual del Directorio estudiantil con búsqueda, filtros, descarga de plantilla, carga masiva, exportación Excel y registro individual dentro del programa asignado.",
        page_break=True,
    )
    evidence_anchor = insert_requirement_evidence(
        evidence_anchor,
        "RF03-A",
        "Clases y matrículas del programa",
        ASSETS / "actual-rf03-clases.jpg",
        "Evidencia RF03-A. Captura actual de la clase con docente, semestre, grupo, ambiente y estudiante matriculado; la asignación determina el alcance visible para Docente y estudiante.",
        page_break=True,
    )
    evidence_anchor = insert_requirement_evidence(
        evidence_anchor,
        "RF03-B",
        "Nuevo módulo de asignación de horarios",
        ASSETS / "actual-rf03-horarios.jpg",
        "Evidencia RF03-B. Nueva pantalla Asignar horarios: Coordinación programa cada clase y la plataforma protege la visibilidad según Docente asignado y matrícula del estudiante.",
        page_break=True,
    )
    evidence_anchor = insert_requirement_evidence(
        evidence_anchor,
        "RF04-A",
        "Registro de asistencia",
        ASSETS / "actual-rf04-asistencia.jpg",
        "Evidencia RF04-A. Captura actual del flujo docente en tres pasos para preparar la lista, marcar cinco estados de asistencia y confirmar el registro por clase y fecha.",
        page_break=True,
    )
    evidence_anchor = insert_requirement_evidence(
        evidence_anchor,
        "RF04-B",
        "Consulta y exportación de reportes",
        ASSETS / "actual-rf04-reportes.jpg",
        "Evidencia RF04-B. Captura actual del reporte docente con filtros, índice de asistencia, gráfico de distribución, actividad por curso y exportación a Excel y PDF.",
        page_break=True,
    )
    evidence_anchor = insert_requirement_evidence(
        evidence_anchor,
        "RF05",
        "Solicitudes y validación de salidas",
        ASSETS / "actual-rf05-permisos.jpg",
        "Evidencia RF05. Captura actual del módulo de Coordinación con contadores para solicitudes pendientes, autorizadas y rechazadas; la autorización genera la boleta oficial que Vigilancia valida mediante QR.",
        page_break=True,
    )
    evidence_anchor = insert_requirement_evidence(
        evidence_anchor,
        "RF06-A",
        "Revisión de justificaciones",
        ASSETS / "actual-rf06-justificaciones.jpg",
        "Evidencia RF06-A. Captura actual de la bandeja de justificaciones con filtros por estado y estudiante; cada decisión registra responsable, fecha y comentario.",
        page_break=True,
    )
    evidence_anchor = insert_requirement_evidence(
        evidence_anchor,
        "RF06-B",
        "Permanencia e incidencias con seguimiento",
        ASSETS / "actual-rf06-permanencia.jpg",
        "Evidencia RF06-B. Captura actual del registro de permanencia con estudiante, responsable, evidencia, historial y acciones para editar, atender o cerrar el caso.",
        page_break=True,
    )

    # 6. Arquitectura y tecnologías
    replace_paragraph(paragraph_after_heading(document, "Arquitectura de la aplicación", 1),
        "La arquitectura separa la presentación, la API y la persistencia. El navegador carga vistas HTML independientes por rol y utiliza JavaScript para consumir la API REST de Express. El servidor valida la sesión, el token CSRF, el rol y el alcance de cada solicitud antes de ejecutar la lógica de negocio y guardar los cambios en SQL Server.")
    architecture_slot = paragraph_by_text(document, "[ INSERTAR DIAGRAMA DE ARQUITECTURA FINAL ]")
    replace_paragraph(architecture_slot, "", style="Body Text")
    architecture_slot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = architecture_slot.add_run().add_picture(str(architecture), width=Inches(7.05))
    add_alt_text(shape, "Arquitectura de SIGA Web", "Diagrama con usuarios, frontend, backend Express y base de datos SQL Server.")
    caption = insert_caption_after(architecture_slot, "Figura 1. Arquitectura final de SIGA Web.")

    tech_rows = [
        ["Editor / IDE", "Visual Studio Code", "Edición, ejecución con F5 y depuración del proyecto."],
        ["Frontend", "HTML5, CSS3 y JavaScript", "Vistas responsive, navegación, formularios e interacción."],
        ["Backend", "Node.js 18+ y Express 5.3", "Servidor web, API REST, sesiones y reglas de negocio."],
        ["Framework / librería", "mssql, ExcelJS, Multer, PDFKit y QRCode", "SQL Server, archivos, exportaciones y códigos QR."],
        ["SGBD", "Microsoft SQL Server 2025 Express", "Persistencia relacional institucional."],
        ["Base de datos", "SIGA", "28 tablas, relaciones, índices y sincronización."],
        ["Servidor / entorno", "Windows, Node.js y SQL Server", "Ejecución local o en intranet; puerto web automático."],
        ["Control de versiones", "No configurado en la carpeta revisada", "La carpeta actual no contiene un repositorio Git activo."],
        ["Hosting / despliegue", "Local / intranet", "Disponible normalmente en http://127.0.0.1:3000."],
    ]
    for i, row in enumerate(tech_rows, 1):
        for j, value in enumerate(row):
            set_cell_text(document.tables[6].rows[i].cells[j], value, size=8.5)
    style_table(document.tables[6], font_size=8.5)
    replace_paragraph(paragraph_after_heading(document, "Tecnologías utilizadas", 1),
        "Las tecnologías se seleccionaron para mantener una instalación sencilla, sin un framework de interfaz, y permitir una integración directa con SQL Server y los formatos de reporte requeridos.")
    clear_paragraph(paragraph_after_heading(document, "Tecnologías utilizadas", 2))

    # 7. Diseño y estructura
    replace_paragraph(paragraph_after_heading(document, "Estructura de carpetas", 1),
        "La organización separa las vistas públicas, la lógica del servidor, los scripts de base de datos, las pruebas y la configuración de ejecución.")
    tree = (
        "SIGA-Web/\n"
        "├── .vscode/             Configuración para iniciar con F5\n"
        "├── database/            Esquema, datos iniciales, sincronización y respaldos\n"
        "├── data/                Persistencia JSON para desarrollo y pruebas\n"
        "├── public/              Interfaz web, estilos, recursos y páginas por rol\n"
        "│   ├── assets/          Logotipo, iconos e imágenes\n"
        "│   └── pages/           Administrador, coordinador, docente, tutor, padre, vigilante y estudiante\n"
        "├── tests/               Pruebas funcionales, permisos y seguridad\n"
        "├── server.js            Servidor Express y API REST\n"
        "├── package.json         Dependencias y scripts de ejecución\n"
        "└── README.md            Instalación, alcance y verificación"
    )
    tree_slot = paragraph_after_heading(document, "Estructura de carpetas", 2)
    replace_paragraph(tree_slot, tree, style="Body Text", size=8.5)
    for run in tree_slot.runs:
        set_run_font(run, name="Consolas", size=8.5)
    for offset in range(3, 9):
        clear_paragraph(paragraph_after_heading(document, "Estructura de carpetas", offset))

    replace_paragraph(paragraph_after_heading(document, "Diseño de interfaces", 1),
        "La interfaz utiliza una estructura común con barra lateral, encabezado contextual, tarjetas de indicadores, filtros, tablas y diálogos. Los componentes conservan estados de foco, etiquetas visibles y tamaños táctiles; en pantallas pequeñas las tablas se transforman en fichas legibles y la navegación lateral funciona como panel móvil.")
    interface_rows = [
        ["Inicio / Login", "Acceso con credenciales institucionales y selección automática del perfil autorizado.", "Acceso implementado en public/index.html"],
        ["Panel principal", "Resumen del rol con indicadores, alertas y accesos directos.", "Captura 1"],
        ["Directorio estudiantil", "Búsqueda, filtros, carga masiva, exportación y registro por programa.", "Captura 2"],
        ["Clases y matrículas", "Asignación de docentes, estudiantes, horarios, aula y modalidad.", "Captura 3"],
        ["Estructura del proyecto", "Organización de archivos en Visual Studio Code.", "Captura 4"],
    ]
    for i, row in enumerate(interface_rows, 1):
        for j, value in enumerate(row):
            set_cell_text(document.tables[7].rows[i].cells[j], value, size=8.5)
    style_table(document.tables[7], font_size=8.5)

    anchor = paragraph_after_heading(document, "Diseño de interfaces", 1)
    anchor = insert_picture_after(anchor, ASSETS / "actual-rf01-coordinador.jpg", 7.15,
                                  "Panel principal de Coordinación", "Panel de coordinación con indicadores y accesos a las tareas del programa.")
    anchor = insert_caption_after(anchor, "Captura 1. Panel principal de Coordinación con indicadores y navegación por tareas.")
    anchor = insert_picture_after(anchor, ASSETS / "actual-rf02-estudiantes.jpg", 7.15,
                                  "Directorio estudiantil", "Directorio de estudiantes con filtros, carga masiva, exportación y acciones.")
    anchor = insert_caption_after(anchor, "Captura 2. Directorio estudiantil con filtros, carga masiva y gestión del registro.")
    anchor = insert_picture_after(anchor, ASSETS / "actual-rf03-clases.jpg", 7.15,
                                  "Clases y matrículas", "Módulo de clases del programa con docente, horario, ambiente y matrícula.")
    anchor = insert_caption_after(anchor, "Captura 3. Módulo de clases, docentes y matrículas del programa de estudios.")
    anchor = insert_picture_after(anchor, ASSETS / "estructura-vscode.png", 2.25,
                                  "Estructura del proyecto en Visual Studio Code", "Carpetas database, public, data y archivos principales del proyecto SIGA Web.")
    insert_caption_after(anchor, "Captura 4. Estructura del proyecto observada en Visual Studio Code.")

    # 8. Componente de datos
    replace_paragraph(paragraph_after_heading(document, "SGBD utilizado", 1),
        "SGBD: Microsoft SQL Server 2025 Express    Base de datos: SIGA    Compatibilidad: 170")
    replace_paragraph(paragraph_after_heading(document, "SGBD utilizado", 2),
        "El modelo relacional contiene 28 tablas. Las claves foráneas vinculan usuarios, roles, programas, estudiantes, clases, matrículas y procesos de seguimiento. La aplicación utiliza una cuenta SQL exclusiva con permisos limitados y confirma cada escritura antes de responder al usuario.", style="List Paragraph")
    erd_slot = paragraph_by_text(document, "[ INSERTAR DIAGRAMA FINAL DE BASE DE DATOS / ENTIDAD-RELACIÓN ]")
    replace_paragraph(erd_slot, "", style="Body Text")
    erd_slot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = erd_slot.add_run().add_picture(str(erd), width=Inches(7.05))
    add_alt_text(shape, "Modelo entidad relación de SIGA", "Diagrama simplificado de programas, usuarios, roles, estudiantes, clases, asistencia, permisos y justificaciones.")
    caption = paragraph_after_heading(document, "[ INSERTAR DIAGRAMA FINAL DE BASE DE DATOS / ENTIDAD-RELACIÓN ]", 1) if False else insert_caption_after(erd_slot, "Figura 2. Modelo entidad relación principal de la base de datos SIGA.")

    replace_paragraph(paragraph_after_heading(document, "Tablas principales", 1),
        "La tabla siguiente resume las entidades que sostienen los procesos académicos y de control de acceso más importantes.")
    db_rows = [
        ["Programas", "Catálogo de programas de estudio.", "ProgramaId, CodigoPublico, Nombre, Estado"],
        ["Usuarios y UsuariosRoles", "Cuentas, credenciales derivadas y permisos multirol.", "UsuarioId, NombreUsuario, ContrasenaHash, RolId, EsPrincipal"],
        ["Estudiantes", "Padrón académico y asignación tutorial.", "EstudianteId, UsuarioId, ProgramaId, TutorUsuarioId, Semestre, Grupo"],
        ["Clases y MatriculasClase", "Oferta académica, docente, horario y estudiantes matriculados.", "ClaseId, DocenteUsuarioId, DiaSemana, Horario, EstudianteId"],
        ["RegistrosAsistencia", "Asistencia diaria y responsable del registro.", "AsistenciaId, ClaseId, EstudianteId, FechaAsistencia, EstadoAsistencia"],
        ["SolicitudesSalida y ValidacionesSalida", "Solicitud, autorización, QR y validación en puerta.", "SolicitudSalidaId, CodigoQr, Estado, VigilanteUsuarioId, FechaValidacion"],
    ]
    for i, row in enumerate(db_rows, 1):
        for j, value in enumerate(row):
            set_cell_text(document.tables[8].rows[i].cells[j], value, size=8.5)
    style_table(document.tables[8], font_size=8.5)

    db_anchor = paragraph_after_heading(document, "Tablas principales", 1)
    db_anchor = insert_picture_after(db_anchor, ASSETS / "sqlserver-siga.png", 4.0,
                                     "Base de datos SIGA en SQL Server Management Studio", "Explorador de objetos de SQL Server con la base de datos SIGA conectada.")
    insert_caption_after(db_anchor, "Captura 5. Base de datos SIGA conectada en SQL Server Management Studio.")

    # Formato final de títulos de las ocho secciones completadas.
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if paragraph.style and paragraph.style.name in ("Heading 1", "Heading 2"):
            if not text:
                paragraph.style = "Body Text"
                continue
            for run in paragraph.runs:
                set_run_font(run, name="Arial", size=(14 if paragraph.style.name == "Heading 1" else 11), bold=True, color=BLACK)

    for table in document.tables:
        mark_table_header(table)

    document.save(str(OUTPUT))
    patch_missing_alt_text(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
