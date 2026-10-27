from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_report import (
    add_alt_text,
    insert_paragraph_after,
    insert_paragraph_after_table,
    insert_requirement_evidence,
    paragraph_by_text,
    patch_missing_alt_text,
    replace_paragraph,
    set_cell_text,
    style_table,
)


ROOT = Path(r"D:\SIGA_WEB\SIGA-Web")
ASSETS = ROOT / "informe_trabajo" / "assets"
SOURCE = ROOT / "Informe_SIGA_hasta_punto_8_con_evidencias.docx"
OUTPUT = ROOT / "Informe_SIGA_hasta_punto_8_evidencias_actualizadas.docx"


def remove_evidence_block(document):
    start = paragraph_by_text(document, "Evidencias visuales de los requerimientos")._p
    end = paragraph_by_text(document, "ARQUITECTURA Y TECNOLOGÍAS")._p
    parent = start.getparent()
    current = start
    while current is not end:
        following = current.getnext()
        parent.remove(current)
        current = following


def replace_picture_before_caption(document, caption_prefix, image_path, title, description, width):
    paragraphs = document.paragraphs
    caption_index = next(index for index, paragraph in enumerate(paragraphs) if paragraph.text.startswith(caption_prefix))
    picture_paragraph = paragraphs[caption_index - 1]
    if not picture_paragraph._p.xpath(".//w:drawing"):
        raise RuntimeError(f"No se encontró la imagen anterior a {caption_prefix}")
    picture_paragraph.clear()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_before = Pt(5)
    picture_paragraph.paragraph_format.space_after = Pt(4)
    shape = picture_paragraph.add_run().add_picture(str(image_path), width=Inches(width))
    add_alt_text(shape, title, description)


def build():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    document = Document(str(SOURCE))
    document.core_properties.title = "Informe final de SIGA Web con evidencias actualizadas"
    document.core_properties.subject = "Secciones 1 a 8 con capturas de la versión vigente al 22 de octubre de 2026"

    evidence_text = {
        "RF01": "Evidencias RF01-A y RF01-B: cambio Coordinador/Docente y menús restringidos.",
        "RF02": "Evidencia RF02: directorio, plantilla, carga masiva, exportación y registro.",
        "RF03": "Evidencias RF03-A y RF03-B: clases, matrículas y nuevo módulo Asignar horarios.",
        "RF04": "Evidencias RF04-A y RF04-B: asistencia en tres pasos y reportes Excel/PDF.",
        "RF05": "Evidencia RF05: revisión de solicitudes y flujo de boleta con validación QR.",
        "RF06": "Evidencias RF06-A y RF06-B: justificaciones y permanencia con seguimiento.",
    }
    requirements_table = document.tables[5]
    for row in requirements_table.rows[1:]:
        code = row.cells[0].text.strip()
        if code in evidence_text:
            set_cell_text(row.cells[4], evidence_text[code], size=8.5)
    style_table(requirements_table, font_size=8.5)

    remove_evidence_block(document)

    anchor = insert_paragraph_after_table(requirements_table)
    replace_paragraph(anchor, "Evidencias visuales de los requerimientos", style="Heading 2", size=14, bold=True)
    anchor.paragraph_format.page_break_before = True
    anchor.paragraph_format.keep_with_next = True
    anchor = insert_paragraph_after(
        anchor,
        "Las capturas fueron obtenidas directamente de la versión vigente de SIGA Web el 22/10/2026. Cada evidencia relaciona el requerimiento con una pantalla funcional y con el alcance visible para la sesión autenticada.",
    )
    anchor.paragraph_format.space_after = Pt(8)

    evidence = [
        ("RF01-A", "Sesión autenticada en modo Coordinador", "actual-rf01-coordinador.jpg", "Evidencia RF01-A. Sesión de Ana Rodríguez Silva en modo Coordinador, con cierre de sesión, selector de rol y acceso limitado al programa Mecatrónica Automotriz."),
        ("RF01-B", "Cambio autorizado al modo Docente", "actual-rf01-docente.jpg", "Evidencia RF01-B. La misma cuenta cambia a modo Docente sin reiniciar sesión y recibe un menú reducido a horario, asistencia, estudiantes, boletas, reportes y perfil."),
        ("RF02", "Gestión de estudiantes y carga masiva", "actual-rf02-estudiantes.jpg", "Evidencia RF02. Directorio estudiantil con búsqueda, filtros, descarga de plantilla, carga masiva, exportación Excel y registro individual dentro del programa asignado."),
        ("RF03-A", "Clases y matrículas del programa", "actual-rf03-clases.jpg", "Evidencia RF03-A. Clase con docente, semestre, grupo, ambiente y estudiante matriculado; la asignación determina el alcance visible para Docente y estudiante."),
        ("RF03-B", "Nuevo módulo de asignación de horarios", "actual-rf03-horarios.jpg", "Evidencia RF03-B. Nueva pantalla Asignar horarios: Coordinación programa cada clase y la plataforma protege la visibilidad según Docente asignado y matrícula del estudiante."),
        ("RF04-A", "Registro de asistencia", "actual-rf04-asistencia.jpg", "Evidencia RF04-A. Flujo docente en tres pasos para preparar la lista, marcar cinco estados de asistencia y confirmar el registro por clase y fecha."),
        ("RF04-B", "Consulta y exportación de reportes", "actual-rf04-reportes.jpg", "Evidencia RF04-B. Reporte docente con filtros, índice de asistencia, gráfico de distribución, actividad por curso y exportación a Excel y PDF."),
        ("RF05", "Solicitudes y validación de salidas", "actual-rf05-permisos.jpg", "Evidencia RF05. Módulo de Coordinación con contadores para solicitudes pendientes, autorizadas y rechazadas; la autorización genera la boleta oficial que Vigilancia valida mediante QR."),
        ("RF06-A", "Revisión de justificaciones", "actual-rf06-justificaciones.jpg", "Evidencia RF06-A. Bandeja de justificaciones con filtros por estado y estudiante; cada decisión registra responsable, fecha y comentario."),
        ("RF06-B", "Permanencia e incidencias con seguimiento", "actual-rf06-permanencia.jpg", "Evidencia RF06-B. Registro de permanencia con estudiante, responsable, evidencia, historial y acciones para editar, atender o cerrar el caso."),
    ]
    for index, (code, title, filename, caption) in enumerate(evidence):
        anchor = insert_requirement_evidence(
            anchor,
            code,
            title,
            ASSETS / filename,
            caption,
            page_break=index > 0,
        )

    replace_picture_before_caption(
        document,
        "Captura 1.",
        ASSETS / "actual-rf01-coordinador.jpg",
        "Panel principal actualizado de Coordinación",
        "Panel principal de Coordinación con menú lateral, selector de rol, cierre de sesión, indicadores y accesos a las tareas del programa.",
        7.15,
    )
    replace_picture_before_caption(
        document,
        "Captura 2.",
        ASSETS / "actual-rf02-estudiantes.jpg",
        "Directorio estudiantil actualizado",
        "Directorio estudiantil con filtros, descarga de plantilla, carga masiva, exportación Excel y gestión del registro.",
        7.15,
    )
    replace_picture_before_caption(
        document,
        "Captura 3.",
        ASSETS / "actual-rf03-clases.jpg",
        "Clases y matrículas actualizadas",
        "Gestión de clases con docente, semestre, grupo, ambiente y estudiantes matriculados.",
        7.15,
    )
    replace_picture_before_caption(
        document,
        "Figura 1.",
        ASSETS / "arquitectura-siga-redisenada.png",
        "Arquitectura final rediseñada de SIGA Web",
        "Diagrama de cuatro capas: usuarios y roles, experiencia web, servicios y seguridad, y persistencia; incluye tecnologías, controles y capacidades transversales.",
        7.15,
    )

    document.save(str(OUTPUT))
    patch_missing_alt_text(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
