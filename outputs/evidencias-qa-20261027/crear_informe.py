from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

p=Path('outputs/evidencias-qa-20261027')
out=p/'Informe de pruebas y evidencias SIGA.docx'
doc=Document()
s=doc.sections[0]; s.top_margin=Inches(.75); s.bottom_margin=Inches(.75); s.left_margin=Inches(.85); s.right_margin=Inches(.85)
styles=doc.styles
for name in ['Normal','Title','Heading 1','Heading 2']:
    styles[name].font.name='Calibri'; styles[name].font.color.rgb=RGBColor.from_string('000000')
styles['Normal'].font.size=Pt(10.5)
styles['Normal'].paragraph_format.space_after=Pt(5)
styles['Title'].font.size=Pt(22)
styles['Heading 1'].font.size=Pt(15)
styles['Heading 2'].font.size=Pt(12)

def para(text, boldlead=None):
    q=doc.add_paragraph()
    if boldlead and text.startswith(boldlead): q.add_run(boldlead).bold=True; q.add_run(text[len(boldlead):])
    else: q.add_run(text)
    return q

def cell_shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.autofit=False
    props=t._tbl.tblPr; borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'5'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    props.append(borders)
    trpr=t.rows[0]._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trpr.append(rep)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.width=Inches(widths[i]); c.text=h; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell_shade(c,'173B69')
        for r in c.paragraphs[0].runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
    for n,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].width=Inches(widths[i]); cells[i].text=val; cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if n%2: cell_shade(cells[i],'F1F5FA')
            for para_ in cells[i].paragraphs:
                para_.paragraph_format.space_before=Pt(3); para_.paragraph_format.space_after=Pt(3)
                for run in para_.runs: run.font.size=Pt(9)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

q=doc.add_paragraph(style='Title'); q.add_run('Informe de pruebas y evidencias de SIGA Web')
para('Pruebas internas realizadas el 27 de octubre de 2026')
para('Resultado: 96 de 96 pruebas automatizadas aprobadas. La revisión visual pasó en los componentes y resoluciones examinados. La aprobación para una carga oficial continúa pendiente de las pruebas y revisiones indicadas en este informe.', 'Resultado:')
para('Las pruebas automatizadas utilizaron datos ficticios y almacenamiento temporal aislado. Las capturas muestran información simulada. La prueba individual y la restauración de SQL Server se observaron en el entorno real, con las limitaciones descritas a continuación.')

doc.add_heading('Resumen de resultados',level=1)
table(['Área','Resultado','Alcance'],[
('Pruebas automatizadas','96 aprobadas','Roles, importación, sesiones, horarios, QR y validaciones cubiertas por la suite.'),
('Piloto Excel ficticio','Aprobado','Diez registros, comparación de campos, duplicados y rechazo de lote mixto.'),
('Revisión visual','Aprobada parcialmente','Acceso, vínculos y matrícula en tamaños seleccionados; capturas con datos ficticios.'),
('Base de datos','Verificada parcialmente','Un registro individual guardado, copia verificada y restaurada en base separada.'),
('UAT y carga oficial','Pendientes','Faltan representantes, conciliación del piloto oficial y aprobación institucional.')],[1.45,1.38,4.0])

doc.add_heading('Pruebas de carga y seguridad',level=1)
para('En el entorno aislado se importaron diez estudiantes ficticios. Se compararon tipo y número de documento, nombre, correo, semestre, grupo, programa y estado. La importación repetida fue rechazada. Un archivo con una fila nueva y otra duplicada también fue rechazado sin guardar la fila nueva.')
para('Una prueba detectó que una forma alternativa de codificar la firma podía producir los mismos bytes. Se reforzó la validación de firmas de sesión y QR para rechazar esa representación y se repitieron las pruebas. Esta observación no demuestra que fuera posible falsificar una firma válida. La corrección se aplica al reiniciar el servidor de SIGA.')

doc.add_heading('Base de datos y recuperación',level=1)
para('En una prueba anterior se creó el estudiante ficticio EST-008 desde la plataforma; se observó en SQL Server y siguió visible tras recargar la página. No se reinició el servidor durante esa prueba y no se guardó una captura de la consulta.')
para('SQL Server Express rechazó la compresión solicitada por el procedimiento de respaldo. Se creó y verificó una copia sin compresión y se restauró en la base separada SIGA_QA_7c294. La comprobación de integridad terminó sin errores observados; se verificaron 29 tablas y el registro de prueba. No se guardó una captura de esa consulta. El script de respaldo corregido pasó una revisión sintáctica, pero no se ejecutó completo como script.')
para('La copia de seguridad y la base restaurada contienen información de la base existente y deben mantenerse con acceso restringido. No se comprobó una copia externa ni una restauración en otro equipo.')

doc.add_heading('Accesibilidad y revisión visual',level=1)
para('Se verificaron el acceso y dos ventanas específicas en 320, 375, 768 y 1366 píxeles con Chrome. Se comprobó la selección de perfil con teclado, la conservación de selecciones y el foco visible de la navegación. El usuario confirmó que Narrador funcionó en la pantalla de acceso; no se grabó audio ni se recorrieron con lector de pantalla todas las secciones y roles.')

doc.add_heading('Pruebas pendientes antes de la carga oficial',level=1)
for item in [
'Pruebas de aceptación con representantes de cada rol y aprobación del responsable institucional.',
'Conciliación de un piloto oficial autorizado contra los archivos fuente antes de la carga completa.',
'Pruebas de interrupción, concurrencia y rendimiento con el volumen esperado de datos.',
'Verificación de accesibilidad por rol y compatibilidad con otros navegadores y dispositivos.',
'Revisión del hosting, HTTPS, correo e integraciones en el entorno real, y verificación de una copia externa.',
'Decidir la conservación o eliminación del registro ficticio, la base restaurada y los archivos de prueba.'
]: doc.add_paragraph(item,style='List Bullet')

doc.add_page_break(); doc.add_heading('Detalle de pruebas automatizadas',level=1)
para('La siguiente relación reproduce los nombres de las pruebas aprobadas en la ejecución registrada. Los resultados corresponden al entorno temporal y no constituyen pruebas manuales de producción.')
lines=(p/'01-pruebas-automatizadas.txt').read_text(encoding='utf-8-sig').splitlines()
entries=[re.sub(r'\s*\([^()]*ms\)\s*$','',line[2:]).strip() for line in lines if line.startswith('✔ ')]
assert len(entries)==96, len(entries)
for i,entry in enumerate(entries,1):
    q=doc.add_paragraph(); q.paragraph_format.space_after=Pt(2); q.paragraph_format.left_indent=Inches(.08); q.paragraph_format.first_line_indent=Inches(-.08)
    r=q.add_run(f'{i:02d}.  '); r.bold=True; q.add_run(entry+' — Aprobada');
    for r in q.runs: r.font.size=Pt(9.5)

doc.add_page_break(); doc.add_heading('Capturas de evidencias visuales',level=1)
para('Las siguientes capturas proceden de pruebas automatizadas con datos ficticios; no representan datos oficiales ni confirman todas las funciones de la plataforma.')
for idx,(file,title,desc,width) in enumerate([
('01-acceso-movil.png','Acceso en vista móvil','Pantalla de acceso con perfiles; ancho de prueba 375 píxeles.',2.0),
('02-vinculos-escritorio.png','Vínculos familiares','Ventana de selección y búsqueda con estudiantes ficticios.',6.1),
('03-matricula-escritorio.png','Matrícula en clases','Ventana de selección de estudiantes ficticios y conservación de la selección.',6.1)]):
    if idx: doc.add_page_break()
    doc.add_heading(title,level=2); para(desc)
    q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.add_run().add_picture(str(p/file),width=Inches(width))
    doc.inline_shapes[-1]._inline.docPr.set('descr',desc)

doc.add_heading('Archivos de respaldo de las evidencias',level=1)
para('El registro completo de ejecución, el resultado visual y las capturas originales se conservan en la carpeta de evidencias del proyecto. El resultado automatizado puede reproducirse con las pruebas del proyecto y almacenamiento temporal aislado.')
doc.save(out)
print(out.resolve())
