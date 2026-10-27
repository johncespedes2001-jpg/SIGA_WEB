from pathlib import Path
import json
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
root=Path('outputs/evidencias-qa-por-pagina')
items=json.loads((root/'resultados-paginas.json').read_text(encoding='utf-8'))
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
for n in ['Normal','Title','Heading 1','Heading 2']:
 s=doc.styles[n]; s.font.name='Calibri'; s.font.color.rgb=RGBColor.from_string('000000')
doc.styles['Normal'].font.size=Pt(10); doc.styles['Normal'].paragraph_format.space_after=Pt(5)
roleNames={'administrador':'Administrador','coordinador':'Coordinador','docente':'Docente','estudiante':'Estudiante','padre':'Padre o apoderado','tutor':'Tutoría del docente','vigilante':'Vigilante'}
doc.add_paragraph('Pruebas y evidencias por página de SIGA Web',style='Title')
doc.add_paragraph('Revisión interna con datos ficticios · 27 de octubre de 2026')
approved=sum(x.get('result')=='aprobada' for x in items)
doc.add_paragraph(f'Se revisaron {len(items)} páginas: {approved} cumplieron las comprobaciones automáticas y {len(items)-approved} requieren revisión. Cada página tiene una ficha con su resultado y, cuando fue posible, su captura.')
doc.add_paragraph('Alcance: carga de cada página con su rol, lectura del título, observación de errores del navegador y respuestas fallidas de API, captura en escritorio y medición del ancho del documento en móvil de 375 píxeles. Las pruebas utilizan el servidor de SIGA y datos temporales ficticios mediante un transporte en memoria; no se modificó SQL Server real.')
doc.add_paragraph('Una página aprobada superó estas comprobaciones iniciales. Esto no certifica que todos sus botones, procesos, accesibilidad, seguridad o datos reales hayan sido probados. Las pruebas manuales con representantes y en dispositivos reales siguen pendientes.')
doc.add_heading('Resultados por perfil',level=1)
for role,label in roleNames.items():
 rows=[x for x in items if x['role']==role]; ok=sum(x.get('result')=='aprobada' for x in rows)
 doc.add_paragraph(f'{label}: {ok} de {len(rows)} páginas aprobadas en la revisión automática.',style='List Bullet')
issues=[x for x in items if x.get('result')!='aprobada']
if issues:
 doc.add_heading('Páginas que requieren revisión',level=1)
 for x in issues:
  d=x.get('desktop',{}); m=x.get('mobile',{}); why=[]
  if x.get('reason'):why.append(x['reason'])
  if x.get('http')!=200:why.append('HTTP '+str(x.get('http')))
  if x.get('url') and x['url']!='/pages/'+x['role']+'/'+x['page']:why.append('Redirigida a '+x['url'])
  if d.get('body'):why.append('Mensaje de error de carga')
  if m.get('width',0)>376:why.append(f"Desbordamiento móvil: {m['width']} px para un ancho de 375 px")
  why += x.get('errors',[])+x.get('apiErrors',[])
  doc.add_paragraph(f"{roleNames.get(x['role'],x['role'])} · {x.get('page','perfil')}: "+('; '.join(why) or 'Resultado por revisar.'),style='List Bullet')
doc.add_heading('Pruebas de registro por tipo de usuario',level=1)
doc.add_paragraph('Las operaciones siguientes se ejecutaron en el servidor de prueba con cuentas y archivos ficticios. Los resultados detallados de las 102 pruebas se incluyen en el anexo. Estas pruebas no modifican usuarios reales ni acreditan que se hayan pulsado todos los botones desde la interfaz.')
for label in ['Administrador','Coordinador','Docente','Padre o apoderado','Vigilante']:
 doc.add_heading(label,level=2)
 doc.add_paragraph('Aprobado: documento inválido y nombre vacío rechazados; creación manual; rechazo de duplicado; edición y lectura del nombre actualizado; desactivación y rechazo de acceso; reactivación e inicio de sesión; eliminación de registro ficticio. Con la plantilla del rol se importó una fila válida, se consultó el rol importado y se rechazó una segunda importación duplicada.')
 if label=='Docente': doc.add_paragraph('Tutoría: se comprobó que la fila con Función de tutor Sí asigna tutoría al docente. No se crea una cuenta independiente de tutor.')
doc.add_heading('Estudiante',level=2)
doc.add_paragraph('Aprobado: carné de extranjería inválido rechazado; creación manual, duplicado rechazado, edición, desactivación, reactivación, inicio de sesión y eliminación del registro ficticio. Además, el piloto Excel de diez estudiantes comprobó los campos importados, duplicados y rechazo de un lote mixto sin guardar parcialmente.')
doc.add_paragraph('Límites: estas pruebas no cubren todas las combinaciones de permisos, todas las contraseñas o archivos posibles, el registro manual desde cada formulario en un navegador real, ni las cargas con información oficial en SQL Server. La validación con representantes y el piloto oficial permanecen pendientes.')
doc.add_page_break()
doc.add_heading('Detalle de la ejecución automatizada',level=1)
import re
log=(root/'resultados-registros-y-roles.txt').read_text(encoding='utf-8-sig').splitlines()
entries=[re.sub(r'\s*\([^()]*ms\)\s*$','',line[2:]).strip() for line in log if line.startswith('✔ ')]
assert len(entries)==102, len(entries)
for i,entry in enumerate(entries,1):
 p=doc.add_paragraph(f'{i:03d}. {entry} — Aprobada'); p.paragraph_format.space_after=Pt(2)
 for run in p.runs: run.font.size=Pt(9)
doc.add_heading('Evidencias individuales por página',level=1)
for x in items:
 doc.add_page_break(); role=x['role']; page=x.get('page','perfil')
 doc.add_heading(roleNames.get(role,role)+' · '+page,level=2)
 doc.add_paragraph('Ruta: /pages/'+role+'/'+page)
 d=x.get('desktop',{}); m=x.get('mobile',{}); result=x.get('result','error')
 doc.add_paragraph('Resultado de carga: '+('Aprobada' if result=='aprobada' else 'Requiere revisión'))
 doc.add_paragraph('Título: '+(d.get('title') or 'No disponible')+' | Ancho en móvil: '+(str(m.get('width'))+' / 375 px' if m.get('width') else 'no disponible'))
 details=[]
 if x.get('reason'):details.append(x['reason'])
 details+=x.get('errors',[])+x.get('apiErrors',[])
 if details: doc.add_paragraph('Observaciones: '+'; '.join(details))
 img=root/'capturas'/(role+'-'+page.replace('.html','')+'.png')
 if img.exists():
  w,h=Image.open(img).size; width=min(6.2,5.6*w/h)
  q=doc.add_paragraph(); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.add_run().add_picture(str(img),width=Inches(width)); doc.inline_shapes[-1]._inline.docPr.set('descr','Captura de '+roleNames.get(role,role)+' '+page+' con datos ficticios')
  doc.add_paragraph('Captura de escritorio con datos ficticios. La medición móvil se realizó por separado.')
 else: doc.add_paragraph('No se obtuvo captura para esta página.')
out=root/'Pruebas completas por rol pagina y registros SIGA.docx'; doc.save(out); print(out.resolve(),len(items),approved)
